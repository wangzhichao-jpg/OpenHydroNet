"""Bounded local CSV/XLSX I/O with explicit units and atomic parameter preview.

Imports update existing IDs only; they never guess graph connectivity or resize pipes.
"""
import base64
import copy
import csv
import io
import json
import math
from zipfile import ZipFile, BadZipFile
from .physics import water, flow_to_si, G
from .validation import validate
from .project import check_project
from .solver import model_hash

COMMON=[('id','稳定 ID',['text']),('name','名称',['text']),('closed','关闭/隔离',['bool'])]
SPEC={
 'node':[('id','稳定 ID',['text']),('name','名称',['text']),('elevation','标高',['m']),('head','定压测压水头 H',['m']),('demand','漏/补水量',['m3/s','m3/h','kg/h','t/h','L/s'])],
 'pipe':COMMON+[('length','实际管长',['m']),('diameter','内径',['m','mm']),('roughness','绝对粗糙度',['m','mm']),('zeta','局阻 ζ',['1']),('local_ratio','沿程附加倍率',['1'])],
 'pump':COMMON+[('speed','相对转速',['1']),('parallel','相同并联台数',['integer']),('efficiency','整机效率',['1'])],
 'valve':COMMON+[('kv','全开 Kv',['m3/h@1bar']),('opening','开度',['1']),('rangeability','可调比',['1'])],
 'terminal':COMMON+[('design_flow','设计流量',['m3/s','m3/h','kg/h','t/h','L/s']),('design_head','设计压降',['m','kPa','Pa'])],
 'resistance':COMMON+[('resistance','等效阻力',['s2/m5'])],
 'flow_control':COMMON+[('set_flow','设定流量',['m3/s','m3/h','kg/h','t/h','L/s']),('min_head','最低资用水头差',['m','kPa','Pa'])],
}


def fields(kind):
    if kind not in SPEC: raise ValueError('不支持的参数表设备类型')
    return [dict(key=k,label=label,units=units,default_unit=units[0]) for k,label,units in SPEC[kind]]


def read_table(payload):
    try: data=base64.b64decode(payload['content_base64'],validate=True)
    except (ValueError,KeyError,TypeError) as exc: raise ValueError('表格文件内容无效') from exc
    if len(data)>3_000_000: raise ValueError('参数表最大 3 MB，请拆分文件')
    kind=payload.get('format');sheet_names=[]
    if kind=='csv':
        encoding=payload.get('encoding','utf-8-sig')
        if encoding not in ('utf-8-sig','gb18030'): raise ValueError('只支持 UTF-8 或 GB18030 CSV')
        try: text=data.decode(encoding)
        except UnicodeError as exc: raise ValueError('CSV 编码不匹配；请选择 UTF-8 或 GB18030 后重新读取') from exc
        delimiter=payload.get('delimiter',',')
        if delimiter not in (',',';','\t'): raise ValueError('不支持的分隔符')
        reader=csv.reader(io.StringIO(text),delimiter=delimiter)
        values=[]
        for row in reader:
            if len(row)>256 or len(values)>10000: raise ValueError('表格限 10000 行、256 列')
            values.append(row)
    elif kind=='xlsx':
        try:
            with ZipFile(io.BytesIO(data)) as z:
                if len(z.infolist())>3000 or sum(x.file_size for x in z.infolist())>40_000_000:
                    raise ValueError('Excel 解压内容过大，拒绝读取')
        except BadZipFile as exc: raise ValueError('不是有效的 XLSX 文件') from exc
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
        from xml.etree.ElementTree import ParseError
        try:book=load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False)
        except (InvalidFileException,ParseError,KeyError,ValueError,TypeError,OSError) as exc:raise ValueError('XLSX 工作簿结构损坏或不受支持') from exc
        try:
            sheet_names=book.sheetnames
            if payload.get('inspect_only'):return dict(sheet_names=sheet_names)
            sheet=payload.get('sheet') or sheet_names[0]
            if sheet not in sheet_names: raise ValueError('工作表不存在')
            ws=book[sheet];values=[]
            for row in ws.iter_rows(values_only=True):
                if len(row)>256 or len(values)>10000: raise ValueError('表格限 10000 行、256 列')
                values.append(list(row))
        finally: book.close()
    else: raise ValueError('仅支持 CSV/XLSX，不支持旧 XLS 或宏工作簿')
    if not values: raise ValueError('表格为空')
    columns=[str(v).strip() if v is not None else '' for v in values[0]]
    while columns and not columns[-1]: columns.pop()
    if not columns or any(not s for s in columns) or len(set(columns))!=len(columns):
        raise ValueError('表头不能为空或重复，请先修正表头')
    rows=[]
    for line,row in enumerate(values[1:],2):
        if all(v in (None,'') for v in row): continue
        if len(row)>len(columns) and any(v not in (None,'') for v in row[len(columns):]):
            raise ValueError(f'第 {line} 行超出表头列数')
        item={key:row[i] if i<len(row) else None for i,key in enumerate(columns)}
        # Dates, error cells, objects and formulas are not evaluated. Strings remain literal.
        if not all(v is None or type(v) in (str,int,float,bool) for v in item.values()):
            raise ValueError(f'第 {line} 行包含日期或不支持的单元格类型')
        rows.append(item)
    return dict(columns=columns,rows=rows,sheet_names=sheet_names,fields={k:fields(k) for k in SPEC},
                notice='只更新已有稳定 ID；ID 列请使用文本，Excel 丢失的前导零不会被猜测恢复；不执行单元格公式')


def _convert(value,unit,rho):
    if unit=='text': return str(value)
    if unit=='bool':
        if type(value) is bool:return value
        text=str(value).strip().lower()
        if text in ('true','1','是','关闭'):return True
        if text in ('false','0','否','开启'):return False
        raise ValueError('关闭状态需为 true/false、1/0、是/否或关闭/开启')
    if type(value) is bool: raise ValueError('布尔值不能作为数值参数')
    try: number=float(value)
    except (ValueError,TypeError) as exc: raise ValueError('不是数字；不执行公式，也不猜测千位/小数分隔符') from exc
    if not math.isfinite(number): raise ValueError('数值必须有限')
    if unit in ('m3/s','m3/h','kg/h','t/h','L/s'):return flow_to_si(number,unit,rho)
    if unit=='mm':return number/1000
    if unit=='kPa':return number*1000/(rho*G)
    if unit=='Pa':return number/(rho*G)
    if unit=='integer':
        if not number.is_integer():raise ValueError('台数必须为整数，不四舍五入')
        return int(number)
    return number


def preview_import(payload):
    project=check_project(payload['project']);kind=payload.get('device_kind','pipe')
    specification={f['key']:f for f in fields(kind)}
    mapping=payload.get('mapping',{});units=payload.get('units',{});rows=payload.get('rows')
    if not isinstance(mapping,dict) or not mapping.get('id') or set(mapping)-specification.keys():
        raise ValueError('必须映射稳定 ID 列，且只能映射当前设备模型字段')
    if not isinstance(rows,list) or not 1<=len(rows)<=10000:raise ValueError('没有可导入的数据行或超过 10000 行')
    model=project['model'];rho,_=water(model.get('fluid',{}).get('temperature',20))
    objects=model['nodes'] if kind=='node' else [e for e in model['edges'] if e['kind']==kind]
    lookup={o['id']:o for o in objects};changes=[];errors=[];seen=set()
    for index,row in enumerate(rows,2):
        if not isinstance(row,dict):raise ValueError('每行必须为列名到单元格的对象')
        raw=row.get(mapping['id']);oid=str(raw).strip() if raw is not None else ''
        if not oid or oid in seen or oid not in lookup:
            errors.append(dict(row=index,id=oid,message='ID 为空、重复或不属于现有指定类型设备；不会自动创建拓扑'));continue
        seen.add(oid);obj=lookup[oid]
        for key,column in mapping.items():
            if key=='id' or not column:continue
            if column not in row:
                errors.append(dict(row=index,id=oid,message=f'映射列 {column} 不存在'));continue
            value=row.get(column)
            if value is None or str(value).strip()=='':
                if not(key=='head' and payload.get('clear_head',False)):continue
                new=None
            else:
                unit=units.get(key,specification[key]['default_unit'])
                if unit not in specification[key]['units']:
                    errors.append(dict(row=index,id=oid,message=f'{key} 的单位不支持'));continue
                try:new=_convert(value,unit,rho)
                except ValueError as exc:errors.append(dict(row=index,id=oid,message=f'{key}: {exc}'));continue
            old=obj.get(key)
            if new!=old:changes.append(dict(row=index,id=oid,field=key,before=old,after=new,si=True));obj[key]=new
    issues=validate(model)
    return dict(project=project,changes=changes,errors=errors,issues=issues,can_apply=not errors and not issues and bool(changes),
                base_project_hash=model_hash(payload['project']),temperature_c=model.get('fluid',{}).get('temperature',20),density_kg_m3=rho,
                note='预览未修改原工程。应用时一次提交，可整体撤销；单位转换已列为 SI，需用户确认。')


def _safe_cell(value):
    if isinstance(value,str) and value.startswith(('=','+','-','@','\t','\r')):return "'"+value
    return value


def workbook_export(project,result=None):
    from openpyxl import Workbook
    book=Workbook();book.remove(book.active)
    for kind,title in [('node','节点'),('pipe','管段'),('pump','水泵'),('valve','阀门'),('terminal','末端'),('resistance','等效阻力'),('flow_control','流控')]:
        ws=book.create_sheet(title);spec=fields(kind);ws.append([f['key'] if f['key'] in ('id','name','closed') else f"{f['key']} [{f['default_unit']}]" for f in spec])
        objects=project['model']['nodes'] if kind=='node' else [e for e in project['model']['edges'] if e['kind']==kind]
        for obj in objects:ws.append([_safe_cell(obj.get(f['key'])) for f in spec])
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    if result:
        from .export import COLUMNS,rows
        ws=book.create_sheet('支路结果');ws.append([label for _,label in COLUMNS])
        for row in rows(project['model'],result):ws.append([_safe_cell(row.get(key)) for key,_ in COLUMNS])
        ws=book.create_sheet('节点结果');keys=['head_m','pressure_pa_gauge','pressure_pa_absolute','boundary_supply_m3_s','mass_residual_m3_s'];ws.append(['id']+keys)
        for oid,row in result['nodes'].items():ws.append([_safe_cell(oid)]+[row.get(k) for k in keys])
    ws=book.create_sheet('计算与假设');ws.append(['字段','值'])
    metadata={'format_version':1,'parameter_units':'参数表默认 SI，表头明确单位；压力/水头不混用','notes':'泵曲线通过工程和设备库随带，参数表不替代完整工程；无效率数据不估算电功率'}
    if result:metadata.update({k:result.get(k) for k in ('version','algorithm','input_hash','calculated_at','residuals','fluid','assumptions','issues')})
    for key,value in metadata.items():ws.append([key,json.dumps(value,ensure_ascii=False) if isinstance(value,(dict,list)) else _safe_cell(value)])
    for ws in book:
        for column in ws.columns:ws.column_dimensions[column[0].column_letter].width=24
    output=io.BytesIO();book.save(output);return output.getvalue()
