import base64
import copy
import io
import json
from pathlib import Path
import pytest
from openpyxl import load_workbook,Workbook
from openhydronet.library import fit_definition,validate_entry
from openhydronet.tableio import read_table,preview_import,workbook_export
from openhydronet.physics import water
from openhydronet import solve

@pytest.fixture
def project():return json.loads((Path(__file__).resolve().parents[1]/'examples/closed-loop.ohn.json').read_text(encoding='utf-8'))
def encoded(data,format='csv',**kwargs):return dict(content_base64=base64.b64encode(data).decode(),format=format,**kwargs)
def table(text,**kwargs):return read_table(encoded(text.encode('utf-8'),**kwargs))

def test_pump_fit_declared_units_and_range():
    r=fit_definition(dict(flows=[0,14.4,28.8,43.2],heads=[32,29.12,20.48,6.08],degree=2,flow_unit='m3/h',temperature=20))
    assert r['curve']==pytest.approx([32,0,-180000],abs=1e-8)
    assert r['valid_flow']==pytest.approx([0,.012]);assert r['rmse_m']<1e-10
    assert r['warnings']==[]
    rho,_=water(80)
    r2=fit_definition(dict(flows=[v*rho for v in [0,14.4,28.8,43.2]],heads=[32,29.12,20.48,6.08],degree=2,flow_unit='kg/h',temperature=80))
    assert r2['curve']==pytest.approx(r['curve'],abs=1e-8)

@pytest.mark.parametrize('payload',[
    dict(flows=[0,0,0],heads=[1,2,3]),dict(flows=[0,1,2],heads=[1,2,3],degree=3),
    dict(flows=[False,1,2],heads=[1,2,3]),dict(flows=[0,1,2],heads=[1,-1,0]),
    dict(flows=[0,1,2],heads=[1,2,3],temperature=True),
])
def test_bad_curve_data_never_generates_fictitious_fit(payload):
    with pytest.raises(ValueError):fit_definition(payload)

def test_library_definition_excludes_runtime_and_topology():
    e=dict(id='L1',kind='pump',name='教学泵',source_kind='synthetic',data_source='虚构教学',parameters=dict(curve=[32,0,-180000],valid_flow=[0,.012]))
    assert validate_entry(e)==e
    for key in ('source','target','id','speed','parallel'):
        bad=copy.deepcopy(e);bad['parameters'][key]=1
        with pytest.raises(ValueError):validate_entry(bad)
    e['data_source']=''
    with pytest.raises(ValueError):validate_entry(e)
    assert validate_entry(e,complete=False)['data_source']==''

def test_utf8_and_gb18030_csv_are_explicit():
    text='编号,内径\nP2,50\n';a=table(text)
    b=read_table(encoded(text.encode('gb18030'),encoding='gb18030'))
    assert a['rows']==b['rows']==[{'编号':'P2','内径':'50'}]
    with pytest.raises(ValueError,match='编码'):read_table(encoded(text.encode('gb18030')))
    with pytest.raises(ValueError,match='表头'):table('id,id\nP2,1\n')


def test_parameter_preview_is_atomic_mm_conversion_and_preserves_ids(project):
    original=copy.deepcopy(project)
    r=preview_import(dict(project=project,device_kind='pipe',rows=[{'ID':'P2','D':'50'}],mapping={'id':'ID','diameter':'D'},units={'diameter':'mm'}))
    assert r['can_apply'];assert r['changes']==[dict(row=2,id='P2',field='diameter',before=.065,after=.05,si=True)]
    assert project==original
    assert r['project']['layout']==project['layout']
    assert [e['id'] for e in r['project']['model']['edges']]==[e['id'] for e in project['model']['edges']]


def test_parameter_mass_flow_uses_current_temperature(project):
    project['model']['fluid']['temperature']=80;rho,_=water(80)
    r=preview_import(dict(project=project,device_kind='terminal',rows=[{'id':'T2','q':str(.006*rho*3600)}],mapping={'id':'id','design_flow':'q'},units={'design_flow':'kg/h'}))
    assert r['can_apply'];assert r['changes'][0]['after']==pytest.approx(.006)
    assert r['density_kg_m3']==rho

@pytest.mark.parametrize('rows,mapping,units',[
    ([{'id':'unknown','v':'50'}],{'id':'id','diameter':'v'},{'diameter':'mm'}),
    ([{'id':'P2','v':'50'},{'id':'P2','v':'40'}],{'id':'id','diameter':'v'},{'diameter':'mm'}),
    ([{'id':'P2','v':'=1+1'}],{'id':'id','diameter':'v'},{'diameter':'mm'}),
    ([{'id':'P2','v':'50'}],{'id':'id','diameter':'missing'},{'diameter':'mm'}),
    ([{'id':'P2','v':'50'}],{'id':'id','diameter':'v'},{'diameter':'not-a-unit'}),
])
def test_invalid_table_never_allows_partial_application(project,rows,mapping,units):
    original=copy.deepcopy(project);r=preview_import(dict(project=project,rows=rows,mapping=mapping,units=units,device_kind='pipe'))
    assert not r['can_apply'];assert r['errors'];assert project==original


def test_integer_count_is_not_rounded(project):
    r=preview_import(dict(project=project,device_kind='pump',rows=[{'id':'PU1','n':'1.5'}],mapping={'id':'id','parallel':'n'},units={'parallel':'integer'}))
    assert not r['can_apply'];assert '不四舍五入' in r['errors'][0]['message']


def test_xlsx_roundtrip_and_formula_safety(project):
    project['model']['edges'][0]['name']='=HYPERLINK("https://invalid.example", "unsafe")'
    result=solve(project['model']);content=workbook_export(project,result)
    book=load_workbook(io.BytesIO(content),data_only=False)
    assert {'节点','管段','水泵','阀门','末端','支路结果','节点结果','计算与假设'}<=set(book.sheetnames)
    assert book['管段']['B2'].value.startswith("'=")
    assert book['管段']['B2'].data_type=='s'
    assert dict(book['计算与假设'].values)['input_hash']==result['input_hash']
    book.close()
    parsed=read_table(encoded(content,'xlsx',sheet='管段'));assert parsed['columns'][3]=='length [m]'
    assert len(parsed['rows'])==4


def test_xlsx_formulas_are_not_evaluated(project):
    b=Workbook();w=b.active;w.append(['id','D']);w.append(['P2','=50+10']);f=io.BytesIO();b.save(f)
    parsed=read_table(encoded(f.getvalue(),'xlsx'))
    assert parsed['rows'][0]['D']=='=50+10'
    preview=preview_import(dict(project=project,device_kind='pipe',rows=parsed['rows'],mapping={'id':'id','diameter':'D'},units={'diameter':'mm'}))
    assert not preview['can_apply'];assert '不执行公式' in preview['errors'][0]['message']
